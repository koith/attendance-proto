#!/usr/bin/env python3
"""Extract the supplied Baekeok Coffee inventory and recipe references.

Usage:
  python tools/build_reference_data_v208.py inventory.xlsx beverage_recipes.pdf food_recipes.pdf output.json
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def menu_name(value: object) -> str:
    value = clean(value)
    value = re.sub(r"\s*/\s*", "/", value)
    value = re.sub(r"\s*\(\s*", " (", value)
    return value.strip()


def inventory_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    section_rows = {cell_range.min_row for cell_range in sheet.merged_cells.ranges if cell_range.min_col == 1 and cell_range.max_col >= 4}
    category = "미분류"
    result = []
    for row in range(3, sheet.max_row + 1):
        name = clean(sheet.cell(row, 1).value)
        if not name:
            continue
        if row in section_rows:
            category = name
            continue
        digest = hashlib.sha1(f"{category}|{name}".encode("utf-8")).hexdigest()[:12].upper()
        result.append({
            "sku": f"SRC-{digest}",
            "name": name,
            "category": category,
            "unit": "기준",
            "minimum_text": clean(sheet.cell(row, 2).value),
            "current_text": clean(sheet.cell(row, 3).value),
            "order_text": clean(sheet.cell(row, 4).value),
            "note": clean(sheet.cell(row, 5).value),
        })
    return result


PAGE_CATEGORIES = {
    1: "커피&콜드브루", 2: "커피&콜드브루",
    3: "신메뉴", 4: "신메뉴", 5: "신메뉴",
    6: "시그니처", 7: "밀크쉐이크", 8: "시그니처",
    9: "라떼&버블티", 10: "라떼&버블티", 11: "라떼&버블티", 12: "라떼&버블티",
    13: "커피&콜드브루", 14: "스무디&에이드", 15: "티&주스",
    16: "스무디&에이드", 17: "스무디&에이드", 18: "스무디&에이드",
    19: "티&주스", 20: "티&주스",
}


def add_recipe(result: list[dict], page_no: int, name: object, variants: list[tuple[object, object]]) -> None:
    name = menu_name(name)
    if not name or name in {"대용량 베이스 사용"}:
        return
    cleaned = []
    for label, content in variants:
        label = clean(label)
        content = str(content or "").strip()
        if label and content:
            cleaned.append({"label": label, "content": content})
    if not cleaned:
        return
    result.append({
        "menu_name": name,
        "category": PAGE_CATEGORIES[page_no],
        "source_page": page_no,
        "source_version": "26.08.21",
        "variants": cleaned,
    })


def beverage_recipe_rows(path: Path) -> list[dict]:
    result: list[dict] = []
    with pdfplumber.open(path) as document:
        for page_no, page in enumerate(document.pages, start=1):
            tables = page.extract_tables()
            if not tables:
                continue
            table = tables[0]
            if len(table) < 3:
                continue
            headers = [clean(value) for value in table[1]]
            width = max(len(row) for row in table)
            for row in table[2:]:
                row = list(row) + [None] * (width - len(row))
                if width == 2:
                    add_recipe(result, page_no, row[0], [(headers[1], row[1])])
                elif width == 3:
                    add_recipe(result, page_no, row[0], list(zip(headers[1:3], row[1:3])))
                elif width == 4:
                    add_recipe(result, page_no, row[0], list(zip(headers[1:4], row[1:4])))
                elif width == 6:
                    add_recipe(result, page_no, row[0], list(zip(headers[1:3], row[1:3])))
                    add_recipe(result, page_no, row[3], list(zip(headers[4:6], row[4:6])))
    return result


def food_recipe_rows(path: Path) -> list[dict]:
    result: list[dict] = []
    with pdfplumber.open(path) as document:
        table = document.pages[1].extract_tables()[0]
        category = "디저트&베이커리"
        category_map = {
            "푸드류 (베이커리)": "디저트&베이커리",
            "푸드류 (백억 휴게소)": "백억휴게소",
            "푸드류 (백억 시네마)": "백억 시네마",
        }
        for row in table[1:]:
            row = list(row) + [None] * (4 - len(row))
            heading = clean(row[0])
            if heading in category_map:
                category = category_map[heading]
                continue
            for name, content in ((row[0], row[1]), (row[2], row[3])):
                name = menu_name(name)
                content = str(content or "").strip()
                if name and content:
                    result.append({
                        "menu_name": name,
                        "category": category,
                        "source_page": 2,
                        "source_version": "26.09.11",
                        "variants": [{"label": "조리·포장", "content": content}],
                    })

        # The latest panel's bulk-base table uses merged cells that split two
        # recipes across rows. Keep these six source recipes explicit so the
        # values remain reviewable instead of relying on a fragile cell merge.
        bases = [
            ("백억 리치 라떼 (15잔)", "①정수 900g\n백억 리치 파우더 600g\n②블렌더 수동 버튼 (20초)"),
            ("복숭아 아이스티 (5잔)", "①아이스티 파우더 300g\n온수 300g\n②바스푼으로 녹이기\n③정수 900g\n다시 섞기"),
            ("초코 베이스 (20잔 분량)", "①온수 400g\n다크 초코 파우더 600g\n카페 시럽 20P(200g)\n②블렌더 수동 버튼 (20초)"),
            ("말차 베이스 (20잔 분량)", "①온수 400g\n말차 파우더 600g\n카페 시럽 40P(400g)\n②블렌더 수동 버튼 (20초)"),
            ("백억커피 (10잔)", "①백억커피 파우더 600g\n카페시럽 10P(100g)\n온수 1000g\n②바스푼으로 녹이기"),
            ("버터 리치 크림 (5잔)", "①컴파운드 휘핑크림 175g\n우유 100g\n리치파우더 30g\n연유파우더 20g\n②블렌딩 4번 버튼"),
        ]
        for name, content in bases:
            result.append({
                "menu_name": name,
                "category": "대용량 베이스",
                "source_page": 1,
                "source_version": "26.09.11",
                "variants": [{"label": "배치 제조", "content": content}],
            })
    return result


def shelf_life_rows(path: Path) -> list[dict]:
    bad = re.compile(r"소세지|전자레인|핫도그")
    result: list[dict] = []
    with pdfplumber.open(path) as document:
        table = document.pages[2].extract_tables()[0]
        categories = ["", ""]
        for row in table[3:]:
            row = list(row) + [None] * (16 - len(row))
            for side, offset in enumerate((0, 8)):
                category = clean(row[offset])
                if category:
                    categories[side] = category
                name = clean(row[offset + 1])
                if not name:
                    continue
                fields = [clean(row[offset + i]) for i in range(2, 8)]
                fields = ["" if bad.search(value) else value for value in fields]
                result.append({
                    "category": categories[side] or "미분류",
                    "name": name,
                    "storage_before": fields[0],
                    "storage_after": fields[1],
                    "expiry_before": fields[2],
                    "expiry_after": fields[3],
                    "after_portion": fields[4],
                    "note": fields[5],
                })
    return result


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("inventory.xlsx beverage_recipes.pdf food_recipes.pdf output.json required")
    inventory_path, beverage_path, food_path, output_path = map(Path, sys.argv[1:])
    shelf_life = shelf_life_rows(food_path)
    shelf_by_name = {clean(row["name"]): row for row in shelf_life}
    inventory = inventory_rows(inventory_path)
    for row in inventory:
        row["shelf_life"] = shelf_by_name.get(clean(row["name"]))
    payload = {
        "version": "v208",
        "source": {
            "inventory": inventory_path.name,
            "beverage_recipes": beverage_path.name,
            "food_recipes": food_path.name,
            "beverage_effective_on": "2026-08-21",
            "food_effective_on": "2026-09-11",
        },
        "inventory": inventory,
        "recipes": beverage_recipe_rows(beverage_path) + food_recipe_rows(food_path),
        "shelf_life": shelf_life,
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "inventory_count": len(payload["inventory"]),
        "recipe_count": len(payload["recipes"]),
        "output": str(output_path),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
